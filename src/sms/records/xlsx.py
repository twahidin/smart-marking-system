"""Render records as markbook.xlsx: sheet Markbook (one row per student, one column per part) and
sheet Rows (every record row flattened, with the student label)."""
import io
from typing import Iterable, List

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from sms.records.builder import Record

REVIEW = "Review"
ROW_HEADERS = ["Student", "Question & part", "Marking scheme answer", "Student's answer (extracted)", "Justification",
               "Awarded mark", "Teacher's mark"]
HEADER_FILL = PatternFill("solid", fgColor="D9E2F3")
REVIEW_FILL = PatternFill("solid", fgColor="FFE8B3")


def _part_labels(records: List[Record]) -> List[str]:
    labels: List[str] = []
    for r in records:
        for row in r.rows:
            if row.label not in labels:
                labels.append(row.label)
    return labels


def _header(ws, headers: List[str]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    ws.freeze_panes = "B2"


def render_xlsx(records: Iterable[Record]) -> bytes:
    records = list(records)
    wb = Workbook()
    mb = wb.active
    mb.title = "Markbook"
    labels = _part_labels(records)
    _header(mb, ["Student"] + labels + ["Total", "Max"])
    for r in records:
        by_label = {row.label: row for row in r.rows}
        cells = [r.student]
        for label in labels:
            row = by_label.get(label)
            cells.append(None if row is None else (REVIEW if row.to_review else row.awarded_marks))
        total = r.total_awarded if r.total_upper == r.total_awarded else r.total_text
        cells += [total, r.total_max]
        mb.append(cells)
        for cell in mb[mb.max_row]:
            if cell.value == REVIEW:
                cell.fill = REVIEW_FILL
    for i in range(1, len(labels) + 4):
        mb.column_dimensions[get_column_letter(i)].width = 22 if i == 1 else 10

    rs = wb.create_sheet("Rows")
    _header(rs, ROW_HEADERS)
    for r in records:
        for row in r.rows:
            rs.append([r.student, row.label, row.scheme_answer, row.student_answer, row.justification, row.awarded,
                       row.teacher or None])
            if row.to_review:
                rs.cell(row=rs.max_row, column=6).fill = REVIEW_FILL
    for i, width in enumerate([22, 14, 40, 40, 40, 16, 14], start=1):
        rs.column_dimensions[get_column_letter(i)].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
