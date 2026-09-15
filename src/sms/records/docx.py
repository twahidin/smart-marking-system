"""Render a Record as a .docx: A4 landscape, header block, the six-column table with shaded header
and amber "Teacher to review" cells, totals rows, and (rubrics) a page with the whole transcription
followed by a page with the bands."""
import io

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from sms.records.builder import Record, TEACHER_TO_REVIEW

HEADERS = ["Question & part", "Marking scheme answer", "Student's answer (extracted)", "Justification",
           "Awarded mark", "Teacher's mark"]
WIDTHS_CM = [2.6, 6.2, 6.2, 6.2, 2.6, 2.6]
HEADER_FILL = "D9E2F3"   # light blue
REVIEW_FILL = "FFE8B3"   # amber
TEACHER_FILL = "F2F2F2"  # grey, the teacher's own column header


def _shade(cell, fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _write(cell, text: str, *, bold: bool = False, size: int = 9) -> None:
    # A fresh cell holds one empty paragraph with no runs; write into it so runs[0] is ours.
    para = cell.paragraphs[0]
    lines = (text or "").split("\n")
    for i, line in enumerate(lines):
        if i:
            para = cell.add_paragraph()
        run = para.add_run(line)
        run.bold = bold
        run.font.size = Pt(size)


def _landscape_a4(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(1.5))


def _when(iso: str) -> str:
    """'2026-09-15T03:06:07Z' -> '2026-09-15 03:06 UTC'; anything else unchanged."""
    if iso and len(iso) == 20 and iso[10] == "T" and iso.endswith("Z"):
        return f"{iso[:10]} {iso[11:16]} UTC"
    return iso or "—"


def _header_block(doc: Document, record: Record) -> None:
    doc.add_heading(record.title, level=1)
    p = doc.add_paragraph()
    p.add_run("Student: ").bold = True
    p.add_run(record.student)
    p = doc.add_paragraph()
    p.add_run("Marked on: ").bold = True
    p.add_run(_when(record.marked_at))
    if record.model:
        p.add_run("    Marked by: ").bold = True
        p.add_run(record.model)
    p = doc.add_paragraph()
    p.add_run("Total: ").bold = True
    p.add_run(f"{record.total_text} / {record.total_max}")
    if record.to_review_count:
        n = record.to_review_count
        p.add_run(f"    {n} part{'s' if n != 1 else ''} to review").bold = True


def _table(doc: Document, record: Record) -> None:
    table = doc.add_table(rows=1, cols=len(HEADERS))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (cell, text) in enumerate(zip(table.rows[0].cells, HEADERS)):
        _write(cell, text, bold=True)
        _shade(cell, TEACHER_FILL if i == len(HEADERS) - 1 else HEADER_FILL)
    for row in record.rows:
        cells = table.add_row().cells
        _write(cells[0], row.label, bold=True)
        _write(cells[1], row.scheme_answer)
        _write(cells[2], row.student_answer)
        _write(cells[3], row.justification)
        _write(cells[4], row.awarded, bold=row.to_review)
        if row.to_review:
            _shade(cells[4], REVIEW_FILL)
        _write(cells[5], row.teacher)
    for label, value, teacher in (("Total awarded", record.total_text, ""),
                                  ("Total max", str(record.total_max), ""),
                                  ("Teacher's total", "", "")):
        cells = table.add_row().cells
        _write(cells[0], label, bold=True)
        _write(cells[4], value, bold=True)
        _write(cells[5], teacher)
    for row in table.rows:
        for cell, width in zip(row.cells, WIDTHS_CM):
            cell.width = Cm(width)


def _transcription_page(doc: Document, record: Record) -> None:
    doc.add_page_break()
    doc.add_heading("Transcription", level=2)
    for para in (record.transcription or "").split("\n\n"):
        p = doc.add_paragraph()
        for i, line in enumerate(para.split("\n")):
            if i:
                p.add_run().add_break()
            p.add_run(line)


def _rubric_page(doc: Document, record: Record) -> None:
    doc.add_page_break()
    doc.add_heading("Rubric", level=2)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ["Criterion", "Band", "Marks", "Descriptor", "Awarded"]):
        _write(cell, text, bold=True)
        _shade(cell, HEADER_FILL)
    for item in record.rubric_page or []:
        bands = item.get("bands") or []
        for b in bands:
            cells = table.add_row().cells
            _write(cells[0], str(item.get("criterion") or ""), bold=True)
            _write(cells[1], str(b.get("band", "")))
            _write(cells[2], str(b.get("marks", "")))
            _write(cells[3], str(b.get("descriptor", "")))
            awarded = item.get("awarded_band") is not None and b.get("band") == item.get("awarded_band")
            _write(cells[4], "✓" if awarded else "", bold=True)
            if awarded:
                _shade(cells[4], HEADER_FILL)


def render_docx(record: Record) -> bytes:
    doc = Document()
    _landscape_a4(doc)
    doc.styles["Normal"].font.size = Pt(10)
    _header_block(doc, record)
    _table(doc, record)
    if record.kind == "rubric" and record.transcription:
        _transcription_page(doc, record)
    if record.kind == "rubric" and record.rubric_page is not None:
        _rubric_page(doc, record)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
