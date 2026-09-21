import io
import zipfile

import openpyxl

from sms.files.render import MAX_DECOMPRESSED, RenderError


def render(data: bytes):
    if not data[:2] == b"PK":
        raise RenderError("workbook", "not an .xlsx workbook")
    # An .xlsx is a zip, and openpyxl inflates it without a budget: check the central directory's
    # declared sizes first, exactly as the .sb3 renderer does, so a 40 KB workbook that claims to
    # expand into gigabytes is refused before a parser ever sees it.
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            declared = sum(i.file_size for i in z.infolist())
    except zipfile.BadZipFile as e:
        raise RenderError("workbook", "not an .xlsx workbook") from e
    if declared > MAX_DECOMPRESSED:
        raise RenderError("workbook", f"workbook expands past {MAX_DECOMPRESSED // (1024 * 1024)} MB")
    try:
        wb_f = openpyxl.load_workbook(io.BytesIO(data), read_only=False, data_only=False, keep_vba=False)
        wb_v = openpyxl.load_workbook(io.BytesIO(data), read_only=False, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise RenderError("workbook", "could not open the workbook") from e
    lines, cells, formulas = [], 0, 0
    for ws in wb_f.worksheets:
        wv = wb_v[ws.title]
        lines.append(f"sheet {ws.title} ({ws.max_row} rows × {ws.max_column} cols)")
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                cells += 1
                v = c.value
                if isinstance(v, str) and v.startswith("="):
                    formulas += 1
                    cached = wv[c.coordinate].value
                    lines.append(f"  {c.coordinate} = {v} → {cached if cached is not None else '(not calculated)'}")
                else:
                    lines.append(f"  {c.coordinate} = {v}")
        if ws.merged_cells.ranges:
            lines.append("  merged: " + ", ".join(str(r) for r in ws.merged_cells.ranges))
        charts = getattr(ws, "_charts", [])
        if charts:
            lines.append(f"  charts: {len(charts)} (" + ", ".join(type(ch).__name__ for ch in charts) + ")")
        dv = getattr(ws.data_validations, "dataValidation", []) if getattr(ws, "data_validations", None) else []
        cf = len(ws.conditional_formatting) if getattr(ws, "conditional_formatting", None) else 0
        if dv or cf:
            lines.append(f"  data validations: {len(dv)}, conditional formats: {cf}")
        lines.append("")
    names = [(n, d.attr_text) for n, d in wb_f.defined_names.items()] if hasattr(wb_f.defined_names, "items") else []
    if names:
        lines.append("named ranges: " + ", ".join(f"{n} = {t}" for n, t in names))
    n_sheets = len(wb_f.worksheets)
    return "\n".join(lines).rstrip() + "\n", f"{n_sheets} sheet{'s' if n_sheets != 1 else ''}, {cells} cells, {formulas} formula{'s' if formulas != 1 else ''}"
