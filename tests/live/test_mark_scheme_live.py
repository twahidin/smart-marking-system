"""Opt-in end-to-end check of the typed-assignment (mark-scheme) pipeline against a real provider:
paper extraction -> scheme extraction -> v2 marking -> marking-record download. Skipped unless
SMS_LIVE_TESTS=1 and a real key is available; prints the provider's error verbatim on failure so a
break is diagnosable without re-running under a debugger."""
import io
import os
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from sms.schemas.scheme import q_label
from sms.web.app import create_app
from sms.web.config import AppConfig
from sms.worker.extract_jobs import run_paper_extract_job, run_scheme_extract_job
from sms.worker.mark_job import run_mark_job

FIXTURES = Path(__file__).parent.parent / "fixtures" / "mark_scheme"

# provider -> env var carrying its key; picks the first one set.
_CANDIDATES = ("tokenrouter", "TOKENROUTER_API_KEY"), ("openrouter", "OPENROUTER_API_KEY"), ("google", "GOOGLE_API_KEY")


def _pick_provider():
    for provider, env_var in _CANDIDATES:
        key = os.environ.get(env_var)
        if key:
            return provider, key
    return None, None


_PROVIDER, _KEY = _pick_provider()

pytestmark = pytest.mark.skipif(
    os.environ.get("SMS_LIVE_TESTS") != "1" or not _PROVIDER,
    reason="set SMS_LIVE_TESTS=1 and one of TOKENROUTER_API_KEY / OPENROUTER_API_KEY / GOOGLE_API_KEY to run",
)


def _read(name: str) -> bytes:
    return (FIXTURES / name).read_bytes()


def test_mark_scheme_assignment_live(tmp_path):
    config = AppConfig(
        database_url=f"sqlite:///{tmp_path / 'live.db'}",
        secret_key="live-test-secret",
        teacher_password="live-test",
        storage_dir=tmp_path / "data",
        embedded_worker=False,
        env={"LLM_PROVIDER": _PROVIDER, "LLM_API_KEY": _KEY},
    )
    app = create_app(config)
    with TestClient(app) as client:
        r = client.post("/api/auth/login", json={"password": "live-test"})
        assert r.status_code == 204

        r = client.post("/api/assignments", json={
            "title": "Live mark-scheme check", "subject": "math", "context": "",
            "rubric": {"criterion_defs": []}, "scheme_kind": "mark_scheme",
        })
        assert r.status_code == 201, r.text
        tid = r.json()["id"]

        r = client.post(f"/api/assignments/{tid}/paper", files=[("files", ("paper.png", _read("paper.png"), "image/png"))])
        assert r.status_code == 200, r.text

        db, storage, settings_store = app.state.db, app.state.storage, app.state.settings_store
        try:
            n_questions = run_paper_extract_job(db, storage, settings_store, tid)
        except Exception as e:  # noqa: BLE001 - surface the provider's own error verbatim
            pytest.fail(f"paper extraction failed against {_PROVIDER}: {e}")
        assert n_questions >= 2

        tpl = next(t for t in client.get("/api/assignments").json() if t["id"] == tid)
        labels = {q_label(q["q_id"]) for q in tpl["questions"]}
        assert {"1(a)", "1(b)"} <= labels, f"expected 1(a) and 1(b) in {labels}"

        r = client.post(f"/api/assignments/{tid}/scheme", files=[("files", ("scheme.png", _read("scheme.png"), "image/png"))])
        assert r.status_code == 200, r.text
        try:
            n_rows = run_scheme_extract_job(db, storage, settings_store, tid)
        except Exception as e:  # noqa: BLE001
            pytest.fail(f"scheme extraction failed against {_PROVIDER}: {e}")
        assert n_rows >= 2

        r = client.post("/api/submissions", data={
            "label": "Live Student", "subject": "math", "context": "", "rubric": '{"criterion_defs": []}',
            "assignment_id": str(tid),
        }, files=[("files", ("script.png", _read("script.png"), "image/png"))])
        assert r.status_code == 202, r.text
        sid = r.json()["id"]

        try:
            run_mark_job(db, storage, settings_store, sid)
        except Exception as e:  # noqa: BLE001
            pytest.fail(f"marking failed against {_PROVIDER}: {e}")

        detail = client.get(f"/api/submissions/{sid}").json()
        assert detail["marks_version"] == 2
        parts = detail["parts"]
        assert len(parts) >= 2
        part_1a = next(p for p in parts if q_label(p["q_id"]) == "1(a)")
        assert any(a["got"] for a in part_1a["awarded"]), part_1a

        r = client.get(f"/api/submissions/{sid}/record.docx")
        assert r.status_code == 200, r.text
        doc = Document(io.BytesIO(r.content))
        text = "\n".join(p.text for p in doc.paragraphs)
        table_text = "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
        assert "1(a)" in text + table_text
