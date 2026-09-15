import io
import zipfile

from openpyxl import load_workbook

from sms.records.builder import Record, RecordRow
from sms.records.bundle import bundle_zip, record_filename
from sms.records.xlsx import render_xlsx


def _row(label, awarded, marks, mx=2, to_review=False):
    return RecordRow(key=label, label=label, scheme_answer="ans", student_answer="x", justification="j", awarded=awarded,
                     teacher="", to_review=to_review, awarded_marks=marks, max_marks=mx)


def _record(student, sid, rows, total, upper, mx, review=0):
    return Record(title="Sec 4 Quadratics", student=student, marked_at="2026-09-15T03:06:07Z", model="m", rows=rows,
                  total_awarded=total, total_upper=upper, total_max=mx, to_review_count=review, rubric_page=None,
                  submission_id=sid, kind="mark_scheme")


REC_A = _record("Tan Wei Ling", 1, [_row("1(a)", "2 / 2", 2), _row("1(b)", "Teacher to review", None, to_review=True)], 2, 4, 4, 1)
REC_B = _record("Lim", 2, [_row("1(a)", "1 / 2", 1), _row("1(b)", "2 / 2 (teacher)", 2), _row("2", "3 / 3", 3, 3)], 6, 6, 7)


def test_markbook_and_rows_sheets():
    data = render_xlsx([REC_A, REC_B])
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Markbook", "Rows"]
    mb = wb["Markbook"]
    rows = [[c.value for c in r] for r in mb.iter_rows()]
    assert rows[0] == ["Student", "1(a)", "1(b)", "2", "Total", "Max"]
    assert rows[1] == ["Tan Wei Ling", 2, "Review", None, "2–4", 4]
    assert rows[2] == ["Lim", 1, 2, 3, 6, 7]
    rs = wb["Rows"]
    rows = [[c.value for c in r] for r in rs.iter_rows()]
    assert rows[0] == ["Student", "Question & part", "Marking scheme answer", "Student's answer (extracted)", "Justification",
                       "Awarded mark", "Teacher's mark"]
    assert len(rows) == 1 + 2 + 3
    assert rows[2] == ["Tan Wei Ling", "1(b)", "ans", "x", "j", "Teacher to review", None]
    assert rows[4] == ["Lim", "1(b)", "ans", "x", "j", "2 / 2 (teacher)", None]


def test_bundle_zip_has_one_docx_per_record_and_the_markbook():
    assert record_filename("Tan Wei Ling") == "tan-wei-ling-marking-record.docx"
    assert record_filename("  Ürün / 3B  ") == "urun-3b-marking-record.docx"
    assert record_filename("***") == "script-marking-record.docx"
    data = bundle_zip([REC_A, REC_B, _record("Lim", 3, [], 0, 0, 0)])
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        names = z.namelist()
        assert names == ["tan-wei-ling-marking-record.docx", "lim-marking-record.docx", "lim-3-marking-record.docx", "markbook.xlsx"]
        assert z.read("tan-wei-ling-marking-record.docx")[:2] == b"PK"
        wb = load_workbook(io.BytesIO(z.read("markbook.xlsx")))
        assert wb["Markbook"].max_row == 4
