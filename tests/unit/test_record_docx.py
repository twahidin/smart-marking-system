import io

from docx import Document
from docx.enum.section import WD_ORIENT

from sms.records.builder import Record, RecordRow
from sms.records.docx import render_docx


def _row(label, awarded, to_review=False, **kw):
    base = dict(key=label, label=label, scheme_answer="x = 3  [M1 A1]", student_answer="x = 3", justification="fine",
                awarded=awarded, teacher="", to_review=to_review, awarded_marks=None if to_review else 2, max_marks=2)
    base.update(kw)
    return RecordRow(**base)


def _record(rubric_page=None, kind="mark_scheme"):
    return Record(title="Sec 4 Quadratics", student="Tan Wei Ling", marked_at="2026-09-15T03:06:07Z", model="openai · gpt-5-mini",
                  rows=[_row("1(a)", "2 / 2"), _row("1(b)", "Teacher to review", to_review=True), _row("2", "2 / 2 (teacher)")],
                  total_awarded=4, total_upper=6, total_max=6, to_review_count=1, rubric_page=rubric_page,
                  submission_id=5, kind=kind)


def _shading(cell):
    tcPr = cell._tc.tcPr
    if tcPr is None:
        return None
    shd = tcPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}shd")
    return None if shd is None else shd.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill")


def test_docx_is_landscape_with_header_block_and_six_column_table():
    data = render_docx(_record())
    assert isinstance(data, bytes) and data[:2] == b"PK"
    doc = Document(io.BytesIO(data))
    section = doc.sections[0]
    assert section.orientation == WD_ORIENT.LANDSCAPE and section.page_width > section.page_height
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Sec 4 Quadratics" in text and "Tan Wei Ling" in text and "2026-09-15" in text and "gpt-5-mini" in text
    assert "4–6 / 6" in text and "1 part to review" in text
    table = doc.tables[0]
    assert len(table.columns) == 6
    header = [c.text for c in table.rows[0].cells]
    assert header == ["Question & part", "Marking scheme answer", "Student's answer (extracted)", "Justification",
                      "Awarded mark", "Teacher's mark"]
    assert all(_shading(c) for c in table.rows[0].cells)
    # 1 header + 3 rows + 3 totals rows (awarded, max, teacher's total)
    assert len(table.rows) == 7
    assert [table.rows[i].cells[0].text for i in (1, 2, 3)] == ["1(a)", "1(b)", "2"]
    assert table.rows[2].cells[4].text == "Teacher to review" and _shading(table.rows[2].cells[4])
    assert table.rows[2].cells[4].paragraphs[0].runs[0].bold
    assert _shading(table.rows[1].cells[4]) is None
    assert table.rows[3].cells[4].text == "2 / 2 (teacher)"
    assert all(table.rows[i].cells[5].text == "" for i in (1, 2, 3))
    assert table.rows[4].cells[0].text == "Total awarded" and table.rows[4].cells[4].text == "4–6"
    assert table.rows[5].cells[0].text == "Total max" and table.rows[5].cells[4].text == "6"
    assert table.rows[6].cells[0].text == "Teacher's total" and table.rows[6].cells[5].text == ""
    assert not doc.inline_shapes  # never any page images


def test_docx_rubric_adds_a_second_page_with_the_bands():
    page = [{"criterion": "Content", "bands": [{"band": "A", "marks": 5, "descriptor": "Rich"}, {"band": "B", "marks": 3, "descriptor": "Some"}],
             "awarded_band": "A"}]
    doc = Document(io.BytesIO(render_docx(_record(rubric_page=page, kind="rubric"))))
    assert len(doc.tables) == 2
    rubric = doc.tables[1]
    assert [c.text for c in rubric.rows[0].cells] == ["Criterion", "Band", "Marks", "Descriptor", "Awarded"]
    assert [c.text for c in rubric.rows[1].cells] == ["Content", "A", "5", "Rich", "✓"]
    assert rubric.rows[2].cells[4].text == ""
    body_xml = doc.element.body.xml
    assert 'w:type="page"' in body_xml  # page break before the rubric
    doc1 = Document(io.BytesIO(render_docx(_record())))
    assert len(doc1.tables) == 1 and 'w:type="page"' not in doc1.element.body.xml


def test_control_characters_are_stripped_by_the_builder_and_render():
    from sms.records.builder import build_record, sanitise
    assert sanitise("a\x00b\x0cc\td\ne\x1f") == "abc\td\ne" and sanitise(None) == ""
    part = {"q_id": "1a", "label": "1(a)", "question_text": "q", "scheme": {"answer": "x\x0b= 3", "marks": [], "notes": ""},
            "extracted": "x\x0c= 3", "workings": "", "illegible": False, "awarded": [], "total": 1, "max": 1,
            "justification": "fine\x0c", "in_scheme": True, "confidence": 0.9, "escalated": False, "reason": None,
            "queue_id": None, "teacher": None}
    d = {"id": 1, "label": "Tan\x00", "context": "", "created_at": "2026-09-15T03:04:05Z", "marks_version": 2, "parts": [part],
         "scheme_kind": "mark_scheme", "assignment_title": "T\x08itle", "totals": {"total": 1, "total_upper": 1, "total_max": 1},
         "marks": [], "rubric": {"criterion_defs": []}, "marked_at": "2026-09-15T03:06:07Z"}
    rec = build_record(d, None, model="m\x0c")
    assert rec.title == "Title" and rec.student == "Tan" and rec.model == "m" and rec.marked_at == "2026-09-15T03:06:07Z"
    row = rec.rows[0]
    assert row.scheme_answer == "x= 3" and row.student_answer == "x= 3" and row.justification == "fine"
    doc = Document(io.BytesIO(render_docx(rec)))
    assert doc.tables[0].rows[1].cells[2].text == "x= 3"
