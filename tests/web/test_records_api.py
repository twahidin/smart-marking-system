import io
import json
import zipfile

from docx import Document
from openpyxl import load_workbook

from tests.web.seed_v2 import seed_v2

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def test_record_requires_auth(client):
    assert client.get("/api/submissions/1/record.docx").status_code == 401
    assert client.post("/api/submissions/records.zip", json={"ids": [1]}).status_code == 401


def test_record_docx_downloads_the_marking_record(auth, app):
    auth.put("/api/settings", json={"provider": "openai", "model": "gpt-5-mini", "api_key": "sk-x", "rpm_limit": 60,
                                    "confidence_threshold": 0})
    tid = auth.post("/api/assignments", json={"title": "Sec 4 Quadratics", "subject": "math", "context": "",
                                              "rubric": {"criterion_defs": [{"id": "c1", "description": "d", "max_score": 1}]},
                                              "scheme_kind": "mark_scheme"}).json()["id"]
    sid, _ = seed_v2(app, label="Tan Wei Ling", assignment_id=tid)
    r = auth.get(f"/api/submissions/{sid}/record.docx")
    assert r.status_code == 200 and r.headers["content-type"].startswith(DOCX)
    assert r.headers["content-disposition"] == 'attachment; filename="tan-wei-ling-marking-record.docx"'
    doc = Document(io.BytesIO(r.content))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Sec 4 Quadratics" in text and "Tan Wei Ling" in text and "gpt-5-mini" in text
    table = doc.tables[0]
    assert [table.rows[i].cells[0].text for i in (1, 2, 3)] == ["1(a)", "1(b)", "2"]
    assert table.rows[3].cells[4].text == "Teacher to review"
    assert table.rows[1].cells[4].text == "2 / 2" and table.rows[1].cells[1].text == "x = 3  [M1 A1]"


def test_record_404_and_409_not_marked(auth, app):
    assert auth.get("/api/submissions/999/record.docx").status_code == 404
    sid = app.state.db.insert("INSERT INTO submissions (label, subject, context, rubric_json, status) "
                              "VALUES ('New', 'math', '', '{\"criterion_defs\": []}', 'queued') RETURNING id")
    r = auth.get(f"/api/submissions/{sid}/record.docx")
    assert r.status_code == 409 and r.json()["error"]["code"] == "not_marked"
    r = auth.post("/api/submissions/records.zip", json={"ids": [sid]})
    assert r.status_code == 409 and r.json()["error"]["code"] == "not_marked"
    r = auth.post("/api/submissions/records.zip", json={"ids": [999]})
    assert r.status_code == 404 and r.json()["error"]["code"] == "not_found"
    assert auth.post("/api/submissions/records.zip", json={"ids": []}).status_code == 400


def test_records_zip_bundles_docx_per_student_and_markbook(auth, app):
    a, _ = seed_v2(app, label="Tan", run_id="ra")
    b, _ = seed_v2(app, label="Lim", run_id="rb", queue={})
    r = auth.post("/api/submissions/records.zip", json={"ids": [a, b]})
    assert r.status_code == 200 and r.headers["content-type"] == "application/zip"
    assert r.headers["content-disposition"] == 'attachment; filename="marking-records.zip"'
    with zipfile.ZipFile(io.BytesIO(r.content)) as z:
        assert z.namelist() == ["tan-marking-record.docx", "lim-marking-record.docx", "markbook.xlsx"]
        wb = load_workbook(io.BytesIO(z.read("markbook.xlsx")))
        rows = [[c.value for c in row] for row in wb["Markbook"].iter_rows()]
        assert rows[0] == ["Student", "1(a)", "1(b)", "2", "Total", "Max"]
        assert rows[1] == ["Tan", 2, 0, "Review", "3–5", 6]
        assert rows[2] == ["Lim", 2, 0, 1, 3, 6]
        assert wb["Rows"].max_row == 1 + 6


def test_v1_record_still_renders(auth, app):
    db = app.state.db
    rubric = {"criterion_defs": [{"id": "c1", "description": "method", "max_score": 2}, {"id": "c2", "description": "answer", "max_score": 3}]}
    sid = db.insert("INSERT INTO submissions (label, subject, context, rubric_json, status, run_id) "
                    "VALUES ('Old', 'math', 'Worksheet 3', :r, 'done', 'r1') RETURNING id", {"r": json.dumps(rubric)})
    final = {"marks": [{"q_id": "q1", "criterion_scores": [2, 3], "total": 5, "confidence": 0.9, "rationale": "r", "evidence": "e"}]}
    db.execute("INSERT INTO marking_runs (run_id, stage, subject, rubric_json, marks_json, final_marks_json, submission_id, final_status) "
               "VALUES ('r1', 'complete', 'math', :rub, :m, :m, :s, 'complete')", {"rub": json.dumps(rubric), "m": json.dumps(final), "s": sid})
    r = auth.get(f"/api/submissions/{sid}/record.docx")
    assert r.status_code == 200
    table = Document(io.BytesIO(r.content)).tables[0]
    assert table.rows[1].cells[0].text == "Q1" and table.rows[1].cells[4].text == "5 / 5"
    assert table.rows[1].cells[1].text == "c1 method (2) · c2 answer (3)"
