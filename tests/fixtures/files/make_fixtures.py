"""Regenerate the file-renderer test fixtures (ok.py, bad_syntax.py, two_sprites.sb3,
broken.sb3, formulas.xlsx).

Run with: uv run python tests/fixtures/files/make_fixtures.py
"""

import io, json, zipfile, pathlib
import openpyxl

HERE = pathlib.Path(__file__).parent

(HERE / "ok.py").write_text("def total(xs):\n    s = 0\n    for x in xs:\n        s += x\n    return s\n\nprint(total([1, 2, 3]))\n")
(HERE / "bad_syntax.py").write_text("def broken(:\n    return 1\n")

project = {"targets": [
  {"isStage": True, "name": "Stage", "variables": {"v1": ["score", 0]}, "lists": {}, "broadcasts": {"b1": "go"}, "blocks": {}, "costumes": [{}], "sounds": []},
  {"isStage": False, "name": "Cat", "variables": {}, "lists": {}, "broadcasts": {}, "costumes": [{}, {}], "sounds": [{}],
   "blocks": {
     "a": {"opcode": "event_whenflagclicked", "next": "b", "parent": None, "inputs": {}, "fields": {}, "topLevel": True},
     "b": {"opcode": "control_repeat", "next": "d", "parent": "a", "inputs": {"TIMES": [1, [6, "10"]], "SUBSTACK": [2, "c"]}, "fields": {}, "topLevel": False},
     "c": {"opcode": "motion_movesteps", "next": None, "parent": "b", "inputs": {"STEPS": [1, [4, "10"]]}, "fields": {}, "topLevel": False},
     "d": {"opcode": "event_broadcast", "next": None, "parent": "b", "inputs": {"BROADCAST_INPUT": [1, [11, "go", "b1"]]}, "fields": {}, "topLevel": False},
   }}]}
buf = io.BytesIO()
with zipfile.ZipFile(buf, "w") as z: z.writestr("project.json", json.dumps(project))
(HERE / "two_sprites.sb3").write_bytes(buf.getvalue())
(HERE / "broken.sb3").write_bytes(b"not a zip at all")

wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Marks"
ws["A1"] = "Item"; ws["B1"] = 10; ws["B2"] = 20; ws["B3"] = 30; ws["B4"] = "=SUM(B1:B3)"
wb.defined_names["Total"] = openpyxl.workbook.defined_name.DefinedName("Total", attr_text="Marks!$B$4")
wb.save(HERE / "formulas.xlsx")
